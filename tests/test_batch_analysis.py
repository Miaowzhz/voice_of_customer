"""验证批量数据从导入、分类到飞书摘要与饼图回传的完整行为。"""

from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
import httpx
from openpyxl import Workbook

from app.api.feishu_webhook import create_app
from app.graph.builder import build_graph
from app.models.classification import FeedbackClassification
from app.repositories.sqlite import SQLiteRepository
from app.runtime import RunService
from app.services.cleaning import load_feedback_file
from app.services.feishu import FeishuAPIError, FeishuClient
from app.services.inputs import bitable_source, decode_feedback_json, normalize_input_rows


def classify(record):
    if "失败" in record["text"]:
        raise RuntimeError("模拟模型超时")
    quality = "粘锅" in record["text"]
    review = "不清楚" in record["text"]
    return FeedbackClassification(
        category="产品质量" if quality else "使用方法与清洁", subcategory="涂层" if quality else "清洁方式",
        sentiment="负面" if quality else "中性", severity="中" if quality else "低",
        sku=record["sku"], is_actionable=quality, evidence=record["sanitized_text"],
        suggested_owner="品控" if quality else "客服", suggested_action="检查涂层" if quality else "补充清洗指南",
        confidence=0.5 if review else 0.95, needs_human_review=review,
    )


class FakeFeishu:
    def __init__(self):
        self.messages = []
        self.images = []
        self.tables = []
        self.table_rows = []

    def send_text(self, receive_id, text, *, receive_id_type):
        self.messages.append((receive_id, text, receive_id_type))

    def upload_image(self, path):
        assert Path(path).read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
        self.uploaded = str(path)
        return "img-report"

    def send_image(self, receive_id, image_key, *, receive_id_type):
        self.images.append((receive_id, image_key, receive_id_type))

    def read_bitable_records(self, source):
        self.tables.append(source)
        return self.table_rows


class InputTests(unittest.TestCase):
    def test_json_envelope_and_chinese_columns(self):
        rows, product = decode_feedback_json('{"product":"锅A","feedbacks":[{"反馈内容":"锅底粘锅"}]}')
        result = normalize_input_rows(rows, product=product)
        self.assertEqual(result[0]["text"], "锅底粘锅")
        self.assertEqual(result[0]["sku"], "锅A")
        self.assertTrue(result[0]["created_at"])

    def test_xlsx_and_json_files_have_same_rows(self):
        with TemporaryDirectory() as directory:
            path = Path(directory)
            workbook = Workbook()
            workbook.active.append(["产品", "反馈内容", "反馈时间"])
            workbook.active.append(["锅A", "锅底粘锅", "2026-09-15"])
            workbook.save(path / "input.xlsx")
            (path / "input.json").write_text(json.dumps([{"产品": "锅A", "反馈内容": "锅底粘锅", "反馈时间": "2026-09-15"}]), encoding="utf-8")
            self.assertEqual(load_feedback_file(path / "input.xlsx"), load_feedback_file(path / "input.json"))

    def test_bitable_rich_text_date_and_record_id(self):
        rows = normalize_input_rows([{"record_id": "rec123", "fields": {
            "产品": "锅A", "反馈内容": [{"text": "锅底"}, {"text": "粘锅"}], "反馈时间": 1789437600000,
        }}])
        self.assertEqual(rows[0]["feedback_id"], "rec123")
        self.assertEqual(rows[0]["text"], "锅底粘锅")
        self.assertRegex(rows[0]["created_at"], r"2026-09-15")

    def test_malformed_and_multi_product_input_is_rejected(self):
        for text in ("{坏 JSON", '42', '{"feedbacks":["无对象"]}'):
            with self.assertRaises(ValueError):
                decode_feedback_json(text)
        with self.assertRaisesRegex(ValueError, "缺少反馈内容列"):
            normalize_input_rows([{"产品": "锅A"}])
        with self.assertRaisesRegex(ValueError, "多个产品"):
            normalize_input_rows([{"text": "a", "sku": "锅A"}, {"text": "b", "sku": "锅B"}])

    def test_offline_mock_keeps_multi_product_import(self):
        path = Path(__file__).resolve().parents[1] / "mock_feedback_50.csv"
        self.assertEqual(len(load_feedback_file(path)), 50)
        with self.assertRaisesRegex(ValueError, "多个产品"):
            load_feedback_file(path, single_product=True)

    def test_bitable_link_requires_table_and_rejects_lookalike_hosts(self):
        source = bitable_source("分析 https://tenant.feishu.cn/base/base123?table=tbl123&view=vew123")
        self.assertEqual(source["view_id"], "vew123")
        self.assertIsNone(bitable_source("https://feishu.cn.attacker.test/base/base123?table=tbl123"))
        with self.assertRaisesRegex(ValueError, "table="):
            bitable_source("https://tenant.feishu.cn/base/base123")


class BatchRuntimeTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.repository = SQLiteRepository(":memory:")
        self.feishu = FakeFeishu()
        self.calls = []

        def classifier(record):
            self.calls.append(record["feedback_id"])
            return classify(record)

        self.service = RunService(
            self.repository, lambda: build_graph(classifier=classifier), feishu_client=self.feishu,
            artifact_root=Path(self.directory.name),
        )

    def tearDown(self):
        self.service.close()
        self.repository.close()
        self.directory.cleanup()

    def test_batch_returns_report_and_chart_despite_review_and_partial_failure(self):
        rows = [
            {"feedback_id": "1", "text": "锅底粘锅"},
            {"feedback_id": "2", "text": "清洗方法不清楚"},
            {"feedback_id": "3", "text": "模拟失败"},
            {"feedback_id": "4", "text": ""},
            {"feedback_id": "1", "text": "重复反馈"},
        ]
        event = {"message_id": "batch1", "chat_id": "chat1", "text": json.dumps({"product": "锅A", "feedbacks": rows}, ensure_ascii=False)}
        run_id = self.service.submit_event(event)
        result = self.service.wait(run_id, timeout=10)
        self.assertEqual(result["status"], "completed")
        self.assertNotIn("__interrupt__", result)
        report = result["analysis_report"]
        self.assertEqual((report["input_count"], report["classified_count"], report["review_count"], report["failure_count"], report["duplicate_count"]), (5, 2, 1, 2, 1))
        self.assertEqual([item["ratio"] for item in report["category_distribution"]], [50.0, 50.0])
        self.assertEqual(len(result["issue_candidates"]), 1)
        self.assertEqual(self.repository.fetch_run(run_id)["source_type"], "batch")
        self.assertEqual(self.repository.fetch_report(run_id)["delivery_status"], "sent")
        self.assertEqual(self.feishu.images, [("chat1", "img-report", "chat_id")])
        self.assertIn("主要问题与建议", self.feishu.messages[-1][1])
        self.assertEqual(self.service.submit_event(event), run_id)
        self.assertEqual(len(self.calls), 3)
        client = TestClient(create_app(repository=self.repository))
        self.assertEqual(client.get(f"/runs/{run_id}/report").json()["report"]["classified_count"], 2)
        self.assertEqual(client.get(f"/runs/{run_id}/chart").headers["content-type"], "image/png")
        self.assertEqual(client.get("/runs/unknown/chart").status_code, 404)

    def test_bitable_message_uses_same_analysis_pipeline(self):
        self.feishu.table_rows = [{"record_id": "rec1", "fields": {"产品": "锅A", "反馈内容": [{"text": "锅底粘锅"}]}}]
        run_id = self.service.submit_event({"message_id": "table1", "chat_id": "chat1", "text": "[产品反馈](https://tenant.feishu.cn/base/base1?table=tbl1&view=vew1)"})
        result = self.service.wait(run_id, timeout=10)
        self.assertEqual(result["analysis_report"]["classified_count"], 1)
        self.assertEqual(self.feishu.tables[0]["view_id"], "vew1")
        self.assertEqual(len(self.feishu.images), 1)

    def test_xlsx_attachment_runs_end_to_end(self):
        path = Path(self.directory.name) / "反馈.xlsx"
        workbook = Workbook()
        workbook.active.append(["SKU", "反馈内容"])
        workbook.active.append(["锅A", "锅底粘锅"])
        workbook.save(path)
        run_id = self.service.submit_event({"message_id": "excel1", "chat_id": "chat1", "file_key": "file1", "local_file_path": str(path)})
        result = self.service.wait(run_id, timeout=10)
        self.assertEqual(result["analysis_report"]["classified_count"], 1)
        self.assertEqual(len(self.feishu.images), 1)

    def test_invalid_json_is_not_silently_classified_as_feedback(self):
        run_id = self.service.submit_event({"message_id": "bad1", "chat_id": "chat1", "text": '{"feedbacks": ['})
        result = self.service.wait(run_id, timeout=5)
        self.assertEqual(result["status"], "failed")
        self.assertIn("JSON 格式错误", self.feishu.messages[-1][1])
        self.assertEqual(self.calls, [])
        self.assertEqual(self.repository.fetch_run(run_id)["status"], "failed")

    def test_image_failure_keeps_analysis_and_marks_delivery(self):
        with patch.object(self.feishu, "upload_image", side_effect=FeishuAPIError("图片权限不足")):
            run_id = self.service.submit_event({"message_id": "imagefail", "chat_id": "chat1", "text": '[{"sku":"锅A","text":"锅底粘锅"}]'})
            result = self.service.wait(run_id, timeout=10)
        self.assertEqual(result["status"], "completed")
        self.assertEqual(self.repository.fetch_report(run_id)["delivery_status"], "partial")
        self.assertIn("饼图发送失败", self.feishu.messages[-1][1])

    def test_concurrent_batches_keep_feedback_and_reports_separate(self):
        text = '[{"feedback_id":"same-id","sku":"锅A","text":"锅底粘锅"}]'
        ids = [self.service.submit_event({"message_id": f"parallel{i}", "chat_id": f"chat{i}", "text": text}) for i in range(2)]
        for run_id in ids:
            self.assertEqual(self.service.wait(run_id, timeout=10)["status"], "completed")
            self.assertIsNotNone(self.repository.fetch_report(run_id))
        rows = self.repository.list_feedback()
        self.assertEqual(len(rows), 2)
        self.assertNotEqual(rows[0]["feedback_id"], rows[1]["feedback_id"])
        self.assertEqual({image[0] for image in self.feishu.images}, {"chat0", "chat1"})


class FeishuBatchTransportTests(unittest.TestCase):
    def test_wiki_resolution_and_permission_error(self):
        def handler(request):
            if "/auth/" in request.url.path:
                return httpx.Response(200, json={"code": 0, "tenant_access_token": "token"})
            if "/get_node" in request.url.path:
                self.assertEqual(request.url.params["token"], "wiki1")
                return httpx.Response(200, json={"code": 0, "data": {"node": {"obj_type": "bitable", "obj_token": "base123"}}})
            self.assertIn("/apps/base123/", request.url.path)
            return httpx.Response(200, json={"code": 91403, "msg": "Forbidden"})

        with httpx.Client(transport=httpx.MockTransport(handler)) as http:
            client = FeishuClient("app", "secret", http_client=http)
            with self.assertRaisesRegex(FeishuAPIError, "91403"):
                client.read_bitable_records({"kind": "wiki", "token": "wiki1", "table_id": "tbl1"})

    def test_expired_token_is_refreshed(self):
        calls = []

        def handler(request):
            calls.append(request)
            return httpx.Response(200, json={"code": 0, "tenant_access_token": f"token{len(calls)}", "expire": 7200})

        with httpx.Client(transport=httpx.MockTransport(handler)) as http:
            client = FeishuClient("app", "secret", http_client=http)
            self.assertEqual(client.tenant_access_token(), "token1")
            self.assertEqual(client.tenant_access_token(), "token1")
            client._token_expires_at = 0
            self.assertEqual(client.tenant_access_token(), "token2")

    def test_paginated_bitable_and_image_upload(self):
        requests = []

        def handler(request):
            requests.append(request)
            if "/auth/" in request.url.path:
                return httpx.Response(200, json={"code": 0, "tenant_access_token": "token", "expire": 7200})
            if "/records" in request.url.path:
                second = request.url.params.get("page_token") == "page2"
                return httpx.Response(200, json={"code": 0, "data": {"items": [{"record_id": "2" if second else "1", "fields": {"反馈内容": "好"}}], "has_more": not second, "page_token": "" if second else "page2"}})
            if request.url.path.endswith("/images"):
                self.assertIn(b'\x89PNG', request.content)
                self.assertIn(b'name="image_type"', request.content)
                return httpx.Response(200, json={"code": 0, "data": {"image_key": "img1"}})
            return httpx.Response(200, json={"code": 0, "data": {"message_id": "om1"}})

        with httpx.Client(transport=httpx.MockTransport(handler)) as http:
            client = FeishuClient("app", "secret", http_client=http)
            rows = client.read_bitable_records({"kind": "base", "token": "base1", "table_id": "tbl1", "view_id": "vew1"})
            self.assertEqual(len(rows), 2)
            pages = [r for r in requests if "/records" in r.url.path]
            self.assertTrue(all(r.url.params["view_id"] == "vew1" for r in pages))
            with TemporaryDirectory() as directory:
                path = Path(directory) / "chart.png"
                path.write_bytes(b"\x89PNG-test")
                image_key = client.upload_image(path)
            client.send_image("chat1", image_key, receive_id_type="chat_id")
            body = json.loads(requests[-1].content)
            self.assertEqual(body["msg_type"], "image")
            self.assertEqual(json.loads(body["content"])["image_key"], "img1")


if __name__ == "__main__":
    unittest.main()
